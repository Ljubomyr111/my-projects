import os, sys, glob
from docx import Document
from pypdf import PdfReader

SRC = r"C:\Users\Lenovo\Downloads\Telegram Desktop"
OUT = r"C:\Users\Lenovo\Documents\deepseek-harness\default-workspace\kursova\src"
os.makedirs(OUT, exist_ok=True)

# 1) plan
plan = os.path.join(SRC, "План курсової роботи.docx")
d = Document(plan)
lines = []
for p in d.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(f"[{p.style.name}] {t}")
for ti, tb in enumerate(d.tables):
    lines.append(f"--- TABLE {ti} ---")
    for row in tb.rows:
        lines.append(" | ".join(c.text.strip() for c in row.cells))
with open(os.path.join(OUT, "plan.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("PLAN paragraphs:", len(d.paragraphs))
print("\n".join(lines[:80]))

# 2) pdfs
for path in sorted(glob.glob(os.path.join(SRC, "*.pdf"))):
    name = os.path.basename(path)
    try:
        r = PdfReader(path)
        n = len(r.pages)
        parts = []
        for i, pg in enumerate(r.pages):
            try:
                parts.append(pg.extract_text() or "")
            except Exception as e:
                parts.append(f"[[extract error p{i}: {e}]]")
        txt = "\n\n".join(parts)
        safe = "".join(ch if ch.isalnum() or ch in " ._-" else "_" for ch in name)[:60]
        with open(os.path.join(OUT, safe + ".txt"), "w", encoding="utf-8") as f:
            f.write(txt)
        print(f"OK  {name}: pages={n} chars={len(txt)}")
    except Exception as e:
        print(f"ERR {name}: {e}")

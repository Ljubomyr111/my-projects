# -*- coding: utf-8 -*-
"""Build the coursework DOCX from text.txt (markers: '# ' chapter, '## ' subsection,
'### ' plain bold heading, '> ' reference-list item, else paragraph)."""
import sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "text.txt")
OUT = os.path.join(HERE, "Rozdil_1_Peredumovy_formuvannia_empirychnoho_pidkhodu.docx")

FONT = "Times New Roman"


def set_run(run, bold=False, italic=False, size=14):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT)


def fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25, spacing=1.5,
            before=0, after=0, keep_next=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.first_line_indent = Cm(indent)
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    return p


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = [l.rstrip("\n") for l in f]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, before=12, after=12,
                    keep_next=True)
            set_run(p.add_run(line[4:]), bold=True)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, before=12, after=12,
                    keep_next=True)
            set_run(p.add_run(line[3:]), bold=True)
        elif line.startswith("# "):
            p = doc.add_paragraph()
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, before=0, after=12,
                    keep_next=True)
            set_run(p.add_run(line[2:]), bold=True)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=0, spacing=1.5)
            p.paragraph_format.left_indent = Cm(1.25)
            set_run(p.add_run(line[2:]), size=14)
        else:
            p = doc.add_paragraph()
            fmt_par(p)
            set_run(p.add_run(line))

    doc.save(OUT)
    body = "\n".join(l for l in lines if l.strip() and not l.startswith(("#", ">")))
    chars = len(body.replace(" ", ""))
    chars_sp = len(body)
    print("saved:", OUT)
    print("paragraphs:", len([l for l in lines if l.strip() and not l.startswith(("#", ">"))]))
    print("chars (no spaces):", chars, "| chars with spaces:", chars_sp)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Convert Markdown to a .docx file via python-docx.

Usage:
  python make_docx.py <input.md> <output.docx>
"""
import sys, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_runs(p, text):
    """Add runs to a paragraph, parsing **bold**, *italic*, `code`."""
    pattern = re.compile(r'(\*\*.+?\*\*|\*.+?\*|`[^`]+`)')
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            p.add_run(part)
    return p

def split_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]

def parse_blocks(md):
    lines = md.split('\n')
    blocks = []
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        if m:
            blocks.append(('heading', len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if re.match(r'^\s*(-{3,}|\*{3,})\s*$', line):
            blocks.append(('hr',))
            i += 1
            continue
        if line.strip().startswith('```'):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(('code', '\n'.join(code)))
            continue
        if '|' in line and i + 1 < n and '-' in lines[i + 1] and re.match(r'^\s*\|?[\s:\-|]+\|?\s*$', lines[i + 1]):
            header = split_row(line)
            i += 2
            rows = []
            while i < n and '|' in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(('table', header, rows))
            continue
        if line.strip().startswith('>'):
            q = []
            while i < n and lines[i].strip().startswith('>'):
                q.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append(('quote', ' '.join(q)))
            continue
        m = re.match(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$', line)
        if m:
            kind = 'ul' if m.group(2) in '-*+' else 'ol'
            items = []
            while i < n:
                mm = re.match(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$', lines[i])
                if not mm:
                    break
                items.append(mm.group(3))
                i += 1
            blocks.append((kind, items))
            continue
        # paragraph
        para = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not re.match(r'^(#{1,6}\s|[-*+]\s|\d+[.)]\s|```|>)', lines[i]) and '|' not in lines[i]:
            para.append(lines[i].strip())
            i += 1
        blocks.append(('para', ' '.join(para)))
    return blocks

def main():
    if len(sys.argv) >= 3:
        md = open(sys.argv[1], encoding='utf-8').read()
        out = sys.argv[2]
    else:
        md = sys.stdin.read()
        out = sys.argv[1] if len(sys.argv) >= 2 else 'output.docx'

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)

    for b in parse_blocks(md):
        kind = b[0]
        if kind == 'heading':
            lvl = b[1]
            p = doc.add_heading(level=min(lvl, 4))
            add_runs(p, b[2])
        elif kind == 'para':
            p = doc.add_paragraph()
            add_runs(p, b[1])
        elif kind == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, '—' * 30)
        elif kind == 'ul':
            for it in b[1]:
                p = doc.add_paragraph(style='List Bullet')
                add_runs(p, it)
        elif kind == 'ol':
            for it in b[1]:
                p = doc.add_paragraph(style='List Number')
                add_runs(p, it)
        elif kind == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            r = p.add_run(b[1])
            r.italic = True
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif kind == 'code':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            for j, cl in enumerate(b[1].split('\n')):
                if j > 0:
                    p.add_run().add_break()
                r = p.add_run(cl)
                r.font.name = 'Consolas'
                r.font.size = Pt(9)
        elif kind == 'table':
            header, rows = b[1], b[2]
            t = doc.add_table(rows=1, cols=len(header))
            t.style = 'Light Grid Accent 1'
            hdr = t.rows[0].cells
            for j, htxt in enumerate(header):
                hdr[j].text = ''
                add_runs(hdr[j].paragraphs[0], htxt)
                for run in hdr[j].paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = t.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ''
                    cells[j].text = ''
                    add_runs(cells[j].paragraphs[0], val)

    doc.save(out)
    print('OK: ' + out)

if __name__ == '__main__':
    main()

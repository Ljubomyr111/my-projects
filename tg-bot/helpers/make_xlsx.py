#!/usr/bin/env python3
"""Convert CSV or JSON (list of lists / list of objects) to .xlsx via openpyxl.

Usage:
  python make_xlsx.py <input.txt> <output.xlsx>
  (input is CSV text, or JSON array)
"""
import sys, json, csv, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

def main():
    if len(sys.argv) >= 3:
        data = open(sys.argv[1], encoding='utf-8').read()
        out = sys.argv[2]
    else:
        data = sys.stdin.read()
        out = sys.argv[1] if len(sys.argv) >= 2 else 'output.xlsx'

    rows = []
    s = data.strip()
    if s.startswith('['):
        obj = json.loads(s)
        if obj and isinstance(obj[0], dict):
            header = list(obj[0].keys())
            rows = [header] + [[r.get(k, '') for k in header] for r in obj]
        else:
            rows = obj
    else:
        rows = list(csv.reader(io.StringIO(data)))

    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)

    if rows:
        # style header
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='305496')
        # freeze header + auto width
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            maxlen = 0
            letter = None
            for c in col:
                if c.value is not None:
                    maxlen = max(maxlen, len(str(c.value)))
                letter = get_column_letter(c.column)
            if letter:
                ws.column_dimensions[letter].width = min(maxlen + 2, 60)

    wb.save(out)
    print('OK: ' + out)

if __name__ == '__main__':
    main()

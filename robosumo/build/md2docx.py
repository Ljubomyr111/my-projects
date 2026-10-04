"""Convert the playbook markdown into a formatted DOCX (then PDF via LibreOffice Kit).

Supports the subset of markdown actually used in playbook.md:
headings, paragraphs, bullet/numbered lists, GFM tables, blockquotes,
horizontal rules, images, and inline **bold** / `code` / [link](url).
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = r"C:\Users\Lenovo\Documents\deepseek-harness\default-workspace\robosumo"
SRC = os.path.join(BASE, "playbook.md")
OUT = os.path.join(BASE, "RoboSumo-Playbook.docx")

INK = RGBColor(0x1C, 0x20, 0x26)
ACCENT = RGBColor(0xC6, 0x28, 0x28)
BLUE = RGBColor(0x15, 0x65, 0xC0)
MUTED = RGBColor(0x60, 0x68, 0x74)
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


# ---------------------------------------------------------------- inline runs
TOKEN = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\((?:[^()]|\([^()]*\))+\))",
    re.S,
)


def add_runs(par, text, base_size=10.5, color=INK):
    """Render inline markdown into a paragraph."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = MONO_FONT
            r.font.size = Pt(base_size - 0.5)
            r.font.color.rgb = RGBColor(0x8B, 0x2A, 0x2A)
        elif part.startswith("[") and "](" in part:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            label, url = m.group(1), m.group(2)
            r = par.add_run(label)
            r.font.color.rgb = BLUE
            r.underline = True
            if not url.startswith("http"):
                r.italic = True
        else:
            r = par.add_run(part)
        r.font.size = Pt(base_size)
        r.font.name = BODY_FONT
        if r.font.color.rgb is None:
            r.font.color.rgb = color
    return par


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def hrule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "C4CCD6")
    bdr.append(bottom)
    pPr.append(bdr)


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.font.color.rgb = INK
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12
    for name, size, color, before, after in [
        ("Heading 1", 20, ACCENT, 16, 6),
        ("Heading 2", 14.5, RGBColor(0x15, 0x20, 0x2C), 14, 5),
        ("Heading 3", 12, RGBColor(0x33, 0x3B, 0x46), 11, 4),
    ]:
        s = doc.styles[name]
        s.font.name = BODY_FONT
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        sec.top_margin = Inches(0.7)
        sec.bottom_margin = Inches(0.7)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        add_runs(p, h, base_size=9.5)
        for r in p.runs:
            r.bold = True
        shade(hdr[i], "EEF2F7")
    for row in body:
        cells = t.add_row().cells
        for i, val in enumerate(row[: len(header)]):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], val, base_size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def main():
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    setup(doc)

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}", stripped):
            hrule(doc)
            i += 1
            continue

        # image
        m = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if m:
            alt, path = m.group(1), m.group(2)
            full = os.path.join(BASE, path.replace("/", os.sep))
            if os.path.exists(full):
                doc.add_picture(full, width=Inches(6.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(alt)
                r.italic = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = MUTED
                r.font.name = BODY_FONT
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 3)
            h = doc.add_heading(level=level)
            add_runs(h, m.group(2))
            for r in h.runs:
                r.font.color.rgb = (
                    ACCENT if level == 1 else
                    RGBColor(0x15, 0x20, 0x2C) if level == 2 else RGBColor(0x33, 0x3B, 0x46))
                r.font.name = BODY_FONT
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(
                r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # code block
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("\n".join(buf))
            r.font.name = MONO_FONT
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x1B, 0x3A, 0x2A)
            pPr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F4F7F5")
            pPr.append(shd)
            continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            add_runs(p, " ".join(buf), base_size=10)
            for r in p.runs:
                r.italic = True
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "C62828")
            bdr.append(left)
            pPr.append(bdr)
            continue

        # lists -- rendered manually so numbering restarts per list and never
        # inherits Word's shared "List Number" counter across sections
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.30 + 0.25 * depth)
            pf.first_line_indent = Inches(-0.18)
            pf.space_after = Pt(2)
            r = p.add_run("\u2022  ")
            r.font.color.rgb = MUTED
            add_runs(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Inches(0.36 + 0.25 * depth)
            pf.first_line_indent = Inches(-0.24)
            pf.space_after = Pt(2)
            r = p.add_run(m.group(2) + ".  ")
            r.bold = True
            r.font.color.rgb = MUTED
            add_runs(p, m.group(3))
            i += 1
            continue

        # paragraph
        par = doc.add_paragraph()
        add_runs(par, stripped)
        i += 1

    doc.save(OUT)
    print("WROTE", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
